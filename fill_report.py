# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = Path("proje_sablonu.docx")
OUT = Path("Veritabani_Proje_Raporu_Doldurulmus_Uysal_Efe_Yuce.docx")
PROJECT = "Spor Salonu Yönetim Sistemi"
STUDENT = "240290054 - Uysal Efe Yüce"


def clear_cell(cell):
    for paragraph in cell.paragraphs:
        paragraph.clear()
    for table in list(cell.tables):
        table._element.getparent().remove(table._element)


def set_cell_text(cell, text):
    clear_cell(cell)
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_border(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_border(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)


def add_heading(cell, text):
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(31, 78, 121)


def add_para(cell, text):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)


def add_bullets(cell, items):
    for item in items:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        run = paragraph.add_run("• " + item)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)


def add_code(cell, text):
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)


def fill_table(table, writer):
    cell = table.cell(0, 0)
    clear_cell(cell)
    set_shading(cell, "F8FBFF")
    set_border(cell)
    writer(cell)


def main():
    doc = Document(SRC)

    meta = doc.tables[0]
    set_cell_text(meta.cell(0, 1), "1")
    set_cell_text(meta.cell(1, 1), STUDENT + " (Takım Lideri)")
    set_cell_text(meta.cell(2, 1), PROJECT)
    style_table(meta)
    for row in meta.rows:
        set_shading(row.cells[0], "EAF2F8")

    def purpose(cell):
        add_para(cell, "Bu projenin amacı, bir spor salonunda yapılan temel yönetim işlemlerini web tabanlı bir sistem üzerinden yönetmektir. Sistem admin, hoca ve üye rollerini ayırarak üyelik takibi, hoca atama, antrenman programı yazma, vücut ölçümü kaydetme, ürün-stok takibi ve satış işlemlerini MySQL veritabanında saklar.")
        add_para(cell, "Uygulama basit ve anlatılabilir bir mimariyle hazırlanmıştır. Önyüz HTML/CSS/JavaScript ile oluşturulmuş, backend Node.js Express ile yazılmış, veriler MySQL tablolarında tutulmuştur.")

    def db_info(cell):
        add_heading(cell, "Veritabanı Yönetim Sistemi")
        add_para(cell, "MySQL kullanılmıştır. Veritabanı adı: spor_salonu")
        add_heading(cell, "Temel Tablolar")
        add_bullets(cell, [
            "uyeler: Üye adı ve telefon numarasını tutar.",
            "personel: Hoca/personel adı ve maaş bilgisini tutar.",
            "uyelik_takibi: Üyelik başlangıç ve bitiş tarihlerini tutar.",
            "antrenman_programi: Hoca tarafından üyeye yazılan program detayını tutar.",
            "vucut_olculeri: Üyenin kilo, boy, yağ oranı ve ölçüm tarihini tutar.",
            "urunler: Ürün adı, kategori, fiyat ve stok miktarını tutar.",
            "satislar: Üyeye yapılan ürün satışlarını, satış tarihini ve adedi tutar.",
        ])
        add_heading(cell, "İlişkiler")
        add_bullets(cell, [
            "uyelik_takibi.uyeid -> uyeler.id",
            "antrenman_programi.uyeid -> uyeler.id",
            "antrenman_programi.personelid -> personel.id",
            "vucut_olculeri.uyeid -> uyeler.id",
            "satislar.uyeid -> uyeler.id",
            "satislar.urunid -> urunler.id",
        ])

    def backend(cell):
        add_para(cell, "Backend JavaScript diliyle Node.js üzerinde geliştirilmiştir. Web framework olarak Express.js kullanılmıştır. MySQL bağlantısı mysql2/promise paketi ile kurulmuştur.")
        add_bullets(cell, [
            "server.js dosyası HTTP API endpointlerini içerir.",
            "dotenv ile .env dosyasından veritabanı bağlantı bilgileri okunur.",
            "cors ve express.json middlewareleri kullanılır.",
            "Admin işlemleri /api/admin/... endpointleriyle; üye ve hoca işlemleri /api/login, /api/programs ve /api/progress endpointleriyle yapılır.",
        ])

    def frontend(cell):
        add_para(cell, "Frontend tarafında HTML, CSS ve sade JavaScript kullanılmıştır. Ek bir framework tercih edilmemiştir; böylece veritabanı ve backend mantığı daha anlaşılır şekilde gösterilmiştir.")
        add_bullets(cell, [
            "index.html: Üye ve hoca giriş ekranı ile panelleri içerir.",
            "admin.html: Admin yönetim panelini içerir.",
            "styles.css: Tüm sayfa tasarımını içerir.",
            "app.js: Üye ve hoca panelindeki API çağrılarını yönetir.",
            "admin.js: Admin form işlemlerini ve stok listesini yönetir.",
        ])

    def architecture(cell):
        add_para(cell, "Proje üç katmanlı basit bir web mimarisiyle çalışır:")
        add_code(cell, """Kullanıcı Tarayıcısı
  |-- admin.html + admin.js  -> Admin işlemleri
  |-- index.html + app.js    -> Üye/Hoca işlemleri
              |
              v
        Express.js Backend (server.js)
              |
              v
          MySQL Veritabanı
  uyeler, personel, uyelik_takibi, antrenman_programi,
  vucut_olculeri, urunler, satislar""")
        add_para(cell, "Admin panelinde yapılan kayıtlar, hoca panelinde yazılan programlar ve üye panelinde görülen bilgiler aynı MySQL veritabanından okunur.")

    def folders(cell):
        add_code(cell, """veritabaniproje/
|-- admin.html        # Admin yönetim paneli
|-- index.html        # Üye/hoca giriş ve panel ekranları
|-- styles.css        # Ortak stil dosyası
|-- app.js            # Üye ve hoca paneli JavaScript işlemleri
|-- admin.js          # Admin paneli JavaScript işlemleri
|-- server.js         # Express backend ve API endpointleri
|-- database.sql      # MySQL tablo oluşturma ve örnek veri scripti
|-- setup-db.js       # database.sql dosyasını MySQL'e kuran script
|-- package.json      # Node.js bağımlılıkları ve komutları
|-- .env              # Veritabanı bağlantı ayarları
|-- baslat.bat        # Kurulum ve backend başlatma kolay komutu""")

    def modules(cell):
        add_bullets(cell, [
            "Admin Modülü: Üye ekleme, personel/hoca ekleme, hoca atama, satış kaydı ve ürün/stok yönetimi işlemlerini yapar.",
            "Üye Modülü: Üyelik bilgilerini, kalan günü, antrenman programını, vücut ölçülerini ve alışverişleri gösterir.",
            "Hoca Modülü: Atanmış üyeleri listeler, üyeye antrenman programı yazar ve gelişim ölçümü ekler.",
            "Stok ve Satış Modülü: Ürünleri listeler, yeni ürün ekler, satışta stok miktarını azaltır.",
            "Veritabanı Kurulum Modülü: database.sql ve setup-db.js ile tabloların kurulmasını sağlar.",
        ])

    for index, writer in enumerate([purpose, db_info, backend, frontend, architecture, folders, modules], start=1):
        fill_table(doc.tables[index], writer)

    team = doc.tables[8]
    team_rows = [["1", "Uysal Efe Yüce", "240290054", "Takım Lideri"]]
    for row_index in range(1, len(team.rows)):
        values = team_rows[row_index - 1] if row_index - 1 < len(team_rows) else ["", "", "", ""]
        for col_index, value in enumerate(values):
            set_cell_text(team.cell(row_index, col_index), value)
    style_table(team)
    for cell in team.rows[0].cells:
        set_shading(cell, "EAF2F8")

    work = doc.tables[9]
    work_rows = [
        ["1", "Veritabanı tasarımı ve tabloların oluşturulması", STUDENT],
        ["2", "Node.js Express backend kurulumu ve MySQL bağlantısı", STUDENT],
        ["3", "Admin paneli: üye, personel, ürün, satış ve atama işlemleri", STUDENT],
        ["4", "Üye ve hoca panelleri: program, ölçüm ve alışveriş akışları", STUDENT],
        ["5", "Test, hata düzeltme ve raporlama", STUDENT],
    ]
    for row_index in range(1, len(work.rows)):
        values = work_rows[row_index - 1] if row_index - 1 < len(work_rows) else ["", "", ""]
        for col_index, value in enumerate(values):
            set_cell_text(work.cell(row_index, col_index), value)
    style_table(work)
    for cell in work.rows[0].cells:
        set_shading(cell, "EAF2F8")

    def screens(cell):
        add_para(cell, "Proje çalıştırıldığında aşağıdaki ekranlar elde edilir:")
        add_bullets(cell, [
            "Ana Giriş Ekranı: Kullanıcı rol seçimi yapılır.",
            "Admin Paneli: Genel özet, üye kaydı, personel kaydı, hoca atama, satış ve stok bölümlerini içerir.",
            "Üye Paneli: Üyelik durumu, aktif program, toplam harcama, ölçüm kayıtları ve alışverişler gösterilir.",
            "Hoca Paneli: Atanmış üyeler listelenir, program ve ölçüm ekleme yapılır.",
            "Stok Ekranı: Ürün listesi, kategori, fiyat, stok ve kritik stok durumu gösterilir.",
        ])
        add_para(cell, "Ekran görüntüleri teslim dosyasına eklenirken ilgili ekranın hangi modüle ait olduğu açıklanmalıdır.")
    fill_table(doc.tables[10], screens)

    def code(cell):
        add_heading(cell, "Backend bağlantısı ve API örneği")
        add_code(cell, """const pool = mysql.createPool({
  host: process.env.MYSQL_HOST,
  user: process.env.MYSQL_USER,
  password: process.env.MYSQL_PASSWORD,
  database: process.env.MYSQL_DATABASE,
});

app.post('/api/admin/members', async (req, res) => {
  const { name, phone, startDate, endDate } = req.body;
  const [result] = await pool.query(
    'INSERT INTO uyeler (ad, telno) VALUES (?, ?)',
    [name, phone]
  );
  await pool.query(
    'INSERT INTO uyelik_takibi (uyeid, baslangictarihi, bitistarihi) VALUES (?, ?, ?)',
    [result.insertId, startDate, endDate]
  );
  res.status(201).json({ ok: true, id: result.insertId });
});""")
        add_heading(cell, "Frontend fetch örneği")
        add_code(cell, """await fetch('/api/admin/members', {
  method: 'POST',
  headers: { 'Content-Type': 'application/json' },
  body: JSON.stringify({ name, phone, startDate, endDate })
});""")
    fill_table(doc.tables[11], code)

    def queries(cell):
        add_heading(cell, "Tablo Oluşturma Örnekleri")
        add_code(cell, """CREATE TABLE uyeler (
  id INT AUTO_INCREMENT PRIMARY KEY,
  ad VARCHAR(100) NOT NULL,
  telno VARCHAR(30) NOT NULL
);

CREATE TABLE antrenman_programi (
  id INT AUTO_INCREMENT PRIMARY KEY,
  uyeid INT NOT NULL,
  personelid INT NOT NULL,
  program_detayi TEXT NOT NULL,
  baslangic_tarihi DATE NOT NULL
);""")
        add_heading(cell, "Listeleme ve Raporlama Sorguları")
        add_code(cell, """SELECT id, ad AS name, telno AS phone
FROM uyeler
ORDER BY id DESC;

SELECT id, urunadi AS name, kategori AS category, fiyat AS price, stokmiktari AS stock
FROM urunler
ORDER BY id;

SELECT s.id, s.satistarihi, ur.urunadi, s.adet, (s.adet * ur.fiyat) AS toplam
FROM satislar s
JOIN urunler ur ON ur.id = s.urunid
WHERE s.uyeid = ?
ORDER BY s.satistarihi DESC;""")
        add_heading(cell, "Kayıt ve Güncelleme Sorguları")
        add_code(cell, """INSERT INTO personel (ad, maas) VALUES (?, ?);
INSERT INTO urunler (urunadi, kategori, fiyat, stokmiktari) VALUES (?, ?, ?, ?);
INSERT INTO vucut_olculeri (uyeid, kilo, boy, yag_orani, olcum_tarihi) VALUES (?, ?, ?, ?, ?);
UPDATE urunler SET stokmiktari = stokmiktari - ? WHERE id = ?;""")
    fill_table(doc.tables[12], queries)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
